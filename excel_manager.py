import os
import io
import csv
from datetime import datetime
from typing import List, Dict, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORTS_DIR = os.path.join(os.path.dirname(__file__), "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)

COLUMNS = [
    {"key": "id", "title": "ID", "width": 8},
    {"key": "name", "title": "Business Name", "width": 32},
    {"key": "distance_text", "title": "Distance", "width": 14},
    {"key": "category", "title": "Category", "width": 22},
    {"key": "rating", "title": "Rating", "width": 10},
    {"key": "reviews", "title": "Reviews", "width": 12},
    {"key": "phone", "title": "Phone Number", "width": 20},
    {"key": "whatsapp", "title": "WhatsApp Chat", "width": 24},
    {"key": "email", "title": "Email Address", "width": 28},
    {"key": "website", "title": "Website", "width": 30},
    {"key": "address", "title": "Address", "width": 40},
    {"key": "instagram", "title": "Instagram", "width": 24},
    {"key": "linkedin", "title": "LinkedIn", "width": 24},
    {"key": "facebook", "title": "Facebook", "width": 24},
    {"key": "maps_url", "title": "Google Maps URL", "width": 35},
    {"key": "status", "title": "Lead Status", "width": 16},
    {"key": "notes", "title": "Notes / Remarks", "width": 30},
]

def generate_excel_workbook(data: List[Dict], sheet_title: str = "Lead Data") -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:30]

    # Modern Palette
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark Navy Slate
    
    regular_font = Font(name="Segoe UI", size=10, color="0F172A")
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    wa_font = Font(name="Segoe UI", size=10, color="059669", underline="single") # Emerald for WhatsApp
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(style='thin', color='E2E8F0')
    border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # 1. Write Header Row
    headers = [col["title"] for col in COLUMNS]
    ws.append(headers)
    ws.row_dimensions[1].height = 30

    for col_idx, col_cfg in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    # 2. Write Data Rows
    for row_idx, item in enumerate(data, 2):
        row_values = []
        for col_cfg in COLUMNS:
            key = col_cfg["key"]
            val = item.get(key, "")
            # Handle distance_text fallback to distance_km if available
            if key == "distance_text" and not val and item.get("distance_km") is not None:
                val = f"{item.get('distance_km')} km"
            row_values.append(val or "")
        
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 24
        
        current_fill = zebra_fill if (row_idx % 2 == 0) else white_fill
        
        for col_idx, col_cfg in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.fill = current_fill
            cell.border = border
            
            # Alignments
            key = col_cfg["key"]
            if key in ["id", "rating", "reviews", "status", "distance_text"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            # Clickable Links
            val_str = str(cell.value or "")
            if key == "whatsapp" and val_str.startswith("http"):
                cell.hyperlink = val_str
                cell.font = wa_font
            elif key in ["website", "maps_url", "instagram", "linkedin", "facebook"] and val_str.startswith("http"):
                cell.hyperlink = val_str
                cell.font = link_font

    # 3. Set Column Widths
    for col_idx, col_cfg in enumerate(COLUMNS, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_cfg["width"]

    # 4. Freeze header row
    ws.freeze_panes = "A2"

    return wb

def export_excel_bytes(data: List[Dict], sheet_title: str = "Lead Data") -> io.BytesIO:
    wb = generate_excel_workbook(data, sheet_title)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def save_excel_file(data: List[Dict], filename: str) -> str:
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb = generate_excel_workbook(data)
    wb.save(filepath)
    return filepath

def export_csv_bytes(data: List[Dict]) -> io.BytesIO:
    output = io.StringIO()
    writer = csv.writer(output)
    
    headers = [col["title"] for col in COLUMNS]
    writer.writerow(headers)
    
    for item in data:
        row = []
        for col in COLUMNS:
            key = col["key"]
            val = item.get(key, "")
            if key == "distance_text" and not val and item.get("distance_km") is not None:
                val = f"{item.get('distance_km')} km"
            row.append(val or "")
        writer.writerow(row)
        
    bytes_output = io.BytesIO()
    # Add UTF-8 BOM so Excel opens Hindi/Special chars correctly
    bytes_output.write(b'\xef\xbb\xbf')
    bytes_output.write(output.getvalue().encode('utf-8'))
    bytes_output.seek(0)
    return bytes_output

def parse_uploaded_file(file_storage) -> List[Dict]:
    filename = file_storage.filename.lower()
    items = []

    if filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        
        raw_headers = [str(h or "").strip().lower() for h in rows[0]]
        
        # Map found headers to keys
        header_map = {}
        for col_cfg in COLUMNS:
            title_lower = col_cfg["title"].lower()
            key = col_cfg["key"]
            for idx, h in enumerate(raw_headers):
                if title_lower in h or key in h:
                    header_map[key] = idx
                    break

        for row_idx, r in enumerate(rows[1:], 1):
            if not any(r):
                continue
            item = {}
            for col_cfg in COLUMNS:
                k = col_cfg["key"]
                idx = header_map.get(k)
                if idx is not None and idx < len(r):
                    val = r[idx]
                    item[k] = str(val).strip() if val is not None else ""
                else:
                    item[k] = ""
            if not item.get("id"):
                item["id"] = row_idx
            items.append(item)

    elif filename.endswith(".csv"):
        stream = io.StringIO(file_storage.stream.read().decode("utf-8-sig", errors="ignore"))
        reader = csv.reader(stream)
        rows = list(reader)
        if not rows:
            return []
        
        raw_headers = [h.strip().lower() for h in rows[0]]
        header_map = {}
        for col_cfg in COLUMNS:
            title_lower = col_cfg["title"].lower()
            key = col_cfg["key"]
            for idx, h in enumerate(raw_headers):
                if title_lower in h or key in h:
                    header_map[key] = idx
                    break

        for row_idx, r in enumerate(rows[1:], 1):
            if not any(r):
                continue
            item = {}
            for col_cfg in COLUMNS:
                k = col_cfg["key"]
                idx = header_map.get(k)
                if idx is not None and idx < len(r):
                    item[k] = r[idx].strip()
                else:
                    item[k] = ""
            if not item.get("id"):
                item["id"] = row_idx
            items.append(item)

    return items
